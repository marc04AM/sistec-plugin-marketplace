"""Merge and validate the per-slice authoring outputs, then write Description/Solution/Note into a
COPY of the Troubleshooting workbook (formatting preserved). Never touches the source file.

Validation: no key authored twice, every workbook row authored, every entry has a description and
a solution, no U+FFFD or control characters. Any failure -> exit code 1 and nothing is saved.

Usage: python fill_excel.py <slices_dir(out_*.json)> <source.xlsx> <dest.xlsx>
                            [--translation <translation.json>]
"""
import argparse
import glob
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font

from common import COL_DESC, COL_NAME, COL_NOTE, COL_SOL, load_json, name_index, row_key

ap = argparse.ArgumentParser()
ap.add_argument('slices')
ap.add_argument('src')
ap.add_argument('dst')
ap.add_argument('--translation')
args = ap.parse_args()

if os.path.abspath(args.src) == os.path.abspath(args.dst):
    sys.exit('ERROR: dest must be a copy, not the source workbook')

BAD = re.compile(r'[�\x00-\x08\x0b\x0c\x0e-\x1f]')
errors = []

authored = {}
for f in sorted(glob.glob(os.path.join(args.slices, 'out_*.json'))):
    d = load_json(f)
    dup = set(d) & set(authored)
    if dup:
        errors.append(f'duplicate keys in {os.path.basename(f)}: {sorted(dup)[:5]}')
    authored.update(d)
print('merged authored keys:', len(authored))
for k, a in authored.items():
    if not (a.get('description') and a.get('solution')):
        errors.append(f'{k}: missing description or solution')
    if any(BAD.search(a.get(c) or '') for c in ('description', 'solution', 'note')):
        errors.append(f'{k}: U+FFFD or control character')

by_name = name_index(args.translation)
wb = openpyxl.load_workbook(args.src)            # keep styles
missing, filled, seen = [], 0, set()
for sh in wb.sheetnames:
    ws = wb[sh]
    hdr = ws.cell(1, COL_NAME)
    nc = ws.cell(1, COL_NOTE, 'Note')
    try:
        nc.font = Font(bold=hdr.font.bold) if hdr.font else Font(bold=True)
    except Exception:
        nc.font = Font(bold=True)
    for r in range(2, ws.max_row + 1):
        key = row_key(ws, r, by_name)
        if not key:
            continue
        seen.add(key)
        a = authored.get(key)
        if not a:
            missing.append(key)
            continue
        ws.cell(r, COL_DESC, a.get('description', ''))
        ws.cell(r, COL_SOL, a.get('solution', ''))
        if a.get('note'):
            ws.cell(r, COL_NOTE, a['note'])
        filled += 1
    for col, w in (('E', 55), ('F', 55), ('G', 30)):
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, w)

if missing:
    errors.append(f'{len(missing)} workbook rows not authored: {missing[:10]}')
extra = set(authored) - seen
if extra:
    print('WARN authored keys not in the workbook (ignored):', len(extra), sorted(extra)[:5])

if errors:
    print('VALIDATION FAILED - nothing saved:')
    for e in errors:
        print('  -', e)
    sys.exit(1)

os.makedirs(os.path.dirname(os.path.abspath(args.dst)), exist_ok=True)
wb.save(args.dst)
print('filled rows:', filled)
print('saved:', args.dst)
