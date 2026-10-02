"""Join every alarm/warning row (from the workbook + alarm_join.json) and split into N balanced
slices for delegated authoring. Rows are ordered by key so same-device instances stay together.
A row with a blank key column is resolved by matching its Name (col D) against the translation.

Usage: python make_slices.py <alarm_join.json> <workbook.xlsx> <out_slices_dir>
                             [--translation <translation.json>] [--n 4]
"""
import argparse
import os

import openpyxl

from common import COL_NAME, dump_json, fb_of, load_json, name_index, row_key

ap = argparse.ArgumentParser()
ap.add_argument('join')
ap.add_argument('xlsx')
ap.add_argument('out')
ap.add_argument('--translation')
ap.add_argument('--n', type=int, default=4)
args = ap.parse_args()

os.makedirs(args.out, exist_ok=True)
j = load_json(args.join)
by_name = name_index(args.translation)
wb = openpyxl.load_workbook(args.xlsx, data_only=True)

recs, unresolved, unknown = [], [], []
for sh in wb.sheetnames:                       # typically 'Alarms', 'Warnings'
    ws = wb[sh]
    for i in range(2, ws.max_row + 1):
        name = ws.cell(i, COL_NAME).value
        k = row_key(ws, i, by_name)
        if not k:
            if name:
                unresolved.append(f'{sh}!{i}: {name}')
            continue
        e = j.get(k)
        if e is None:
            unknown.append(k)
            e = {}
        recs.append({
            'key': k, 'sheet': sh,
            'name': name,                       # operator-facing text = meaning source of truth
            'meaning': e.get('meaning', ''),
            'italian': e.get('italian', ''),
            'condition': e.get('condition', ''),
            'coded': e.get('coded'),
            'fb': fb_of(k),
        })

recs.sort(key=lambda r: (r['sheet'], r['key']))
size = -(-len(recs) // args.n) if recs else 0
for idx in range(args.n):
    chunk = recs[idx * size:(idx + 1) * size] if size else []
    if not chunk:
        continue
    dump_json(chunk, os.path.join(args.out, f'slice_{idx+1}.json'))
    print(f'slice_{idx+1}: {len(chunk)} rows  [{chunk[0]["key"]} .. {chunk[-1]["key"]}]')
print('total:', len(recs))
if unresolved:
    print('WARN rows without a key that could not be resolved by Name:', len(unresolved), unresolved[:5])
if unknown:
    print('WARN keys not defined in the PLC (no join entry):', len(unknown), unknown[:5])
