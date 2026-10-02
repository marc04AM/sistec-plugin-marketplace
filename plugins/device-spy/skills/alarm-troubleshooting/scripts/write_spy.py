"""Write the --spy dataset (spy_all.json + authored spy_slices/out_*.json) to an output file.
The extension picks the format: .xlsx (Troubleshooting sheet look), .csv (UTF-8 BOM), .json, or .md.
Columns: priority, key, area, it, en, description, solution, note.

Usage: python write_spy.py <spy_all.json> <spy_slices_dir> <output_path>
"""
import csv
import glob
import json
import os
import sys

from common import load_json

SPY_ALL, SLICES, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
COLS = ['priority', 'key', 'area', 'it', 'en', 'description', 'solution', 'note']

rows = load_json(SPY_ALL)
authored = {}
for f in sorted(glob.glob(os.path.join(SLICES, 'out_*.json'))):
    authored.update(load_json(f))

table = []
for k, r in rows.items():
    a = authored.get(k, {})
    table.append({
        'priority': r['priority'], 'key': r['key'], 'area': r['area'],
        'it': r['it'], 'en': r['en'],
        'description': r['description'] if r['description'] is not None else a.get('description', ''),
        'solution': r['solution'] if r['solution'] is not None else a.get('solution', ''),
        'note': (r['note'] if r['note'] is not None else a.get('note', '')) or '',
    })
# order: alarms first, then warnings; within, by key
table.sort(key=lambda x: (0 if x['priority'] == 'Alarm' else 1, x['key']))
blanks = [x['key'] for x in table if not x['description'] or not x['solution']]

ext = os.path.splitext(OUT)[1].lower()
os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)

if ext in ('.xlsx', '.xlsm'):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Troubleshooting'
    ws.append([c.upper() if c in ('it', 'en') else c.capitalize() for c in COLS])
    hf = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='305496')
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(1, c)
        cell.font = hf
        cell.fill = fill
        cell.alignment = Alignment(vertical='center')
    for x in table:
        ws.append([x[c] for c in COLS])
    widths = {'priority': 10, 'key': 34, 'area': 8, 'it': 45, 'en': 45, 'description': 60, 'solution': 60, 'note': 30}
    from openpyxl.utils import get_column_letter
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[c]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'
    wb.save(OUT)
elif ext == '.csv':
    with open(OUT, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction='ignore')
        w.writeheader()
        w.writerows(table)
elif ext == '.json':
    json.dump(table, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
elif ext in ('.md', '.markdown'):
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('| ' + ' | '.join(COLS) + ' |\n')
        f.write('|' + '|'.join(['---'] * len(COLS)) + '|\n')
        for x in table:
            f.write('| ' + ' | '.join(str(x[c]).replace('|', '\\|').replace('\n', ' ') for c in COLS) + ' |\n')
else:
    print('ERROR unsupported extension:', ext)
    sys.exit(2)

print('rows written:', len(table), '| alarms:', sum(1 for x in table if x['priority'] == 'Alarm'),
      '| warnings:', sum(1 for x in table if x['priority'] == 'Warning'))
if blanks:
    print('WARN rows missing description/solution:', len(blanks), blanks[:8])
print('saved:', OUT)
